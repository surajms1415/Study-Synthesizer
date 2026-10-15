import os
import uuid
import asyncio
from typing import List
from fastapi import UploadFile
from models import DocumentContent
import PyPDF2

def extract_text_from_pdf(file_obj, filename: str) -> List[DocumentContent]:
    docs = []
    reader = PyPDF2.PdfReader(file_obj)
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text:
            docs.append(DocumentContent(
                document_id=str(uuid.uuid4()),
                document_name=filename,
                source_type='pdf',
                text=text,
                metadata={'page_number': i + 1}
            ))
    return docs

def extract_content_from_pptx(file_obj, filename: str) -> List[DocumentContent]:
    from pptx import Presentation
    prs = Presentation(file_obj)
    docs = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        text = "\n".join(slide_text)
        if text.strip():
            docs.append(DocumentContent(
                document_id=str(uuid.uuid4()),
                document_name=filename,
                source_type='pptx',
                text=text,
                metadata={'slide_number': i + 1}
            ))
    return docs

def extract_content_from_docx(file_obj, filename: str) -> List[DocumentContent]:
    from docx import Document
    doc = Document(file_obj)
    text = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(text)
    if full_text.strip():
        return [DocumentContent(
            document_id=str(uuid.uuid4()),
            document_name=filename,
            source_type='docx',
            text=full_text,
            metadata={}
        )]
    return []

async def ingest_document(upload_file: UploadFile) -> List[DocumentContent]:
    ext = os.path.splitext(upload_file.filename)[1].lower()
    file_obj = upload_file.file
    filename = upload_file.filename
    if ext == '.pdf':
        return extract_text_from_pdf(file_obj, filename)
    elif ext in ['.pptx', '.ppt']:
        return extract_content_from_pptx(file_obj, filename)
    elif ext in ['.docx', '.doc']:
        return extract_content_from_docx(file_obj, filename)
    elif ext == '.txt':
        text = (await upload_file.read()).decode('utf-8')
        return [DocumentContent(
            document_id=str(uuid.uuid4()),
            document_name=filename,
            source_type='txt',
            text=text,
            metadata={}
        )]
    return []

def _fetch_transcript_sync(video_id: str):
    from youtube_transcript_api import YouTubeTranscriptApi
    
    # 1. Try instance .fetch() (youtube_transcript_api >= 1.0)
    try:
        api = YouTubeTranscriptApi()
        if hasattr(api, 'fetch'):
            res = api.fetch(video_id)
            if hasattr(res, 'to_raw_data'):
                return res.to_raw_data()
            return res
    except Exception as e:
        # If default language fails, try listing available transcripts
        try:
            api = YouTubeTranscriptApi()
            if hasattr(api, 'list'):
                t_list = api.list(video_id)
                for t in t_list:
                    fetched = t.fetch()
                    if hasattr(fetched, 'to_raw_data'):
                        return fetched.to_raw_data()
                    return fetched
        except Exception:
            pass

    # 2. Try static .get_transcript() (older youtube_transcript_api)
    if hasattr(YouTubeTranscriptApi, 'get_transcript'):
        return YouTubeTranscriptApi.get_transcript(video_id)

    raise RuntimeError("No transcript available or method failed")

async def ingest_youtube_transcript(yt_url: str) -> List[DocumentContent]:
    import re
    
    match = re.search(r"(?:v=|\/)([0-9A-Za-z_-]{11})(?:\?|&|$)", yt_url)
    if not match: match = re.search(r"youtu\.be\/([0-9A-Za-z_-]{11})(?:\?|&|$)", yt_url)
    video_id = match.group(1) if match else None
    
    if not video_id:
        return []
        
    try:
        transcript_list = await asyncio.to_thread(_fetch_transcript_sync, video_id)
        docs = []
        for segment in transcript_list:
            text = segment.get('text', '') if isinstance(segment, dict) else getattr(segment, 'text', '')
            start = segment.get('start', 0.0) if isinstance(segment, dict) else getattr(segment, 'start', 0.0)
            duration = segment.get('duration', 0.0) if isinstance(segment, dict) else getattr(segment, 'duration', 0.0)
            if text and text.strip():
                docs.append(DocumentContent(
                    document_id=str(uuid.uuid4()),
                    document_name=f"YouTube: {video_id}",
                    source_type='youtube',
                    text=text,
                    metadata={
                        'start_time': start,
                        'end_time': start + duration,
                        'source_url': yt_url
                    }
                ))
        return docs
    except Exception as e:
        print(f"No yt transcript: {e}")
        return []
